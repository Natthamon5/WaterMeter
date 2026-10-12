import gspread
from dotenv import load_dotenv
from flask import Flask, render_template, request, jsonify, session, send_file
import os, json, io, csv, base64
from datetime import datetime
import prediction
from openpyxl import Workbook
from openpyxl.drawing.image import Image as ExcelImage

app = Flask(__name__)
app.secret_key = "supersecret" 

load_dotenv()
SHEET_ID = os.getenv("SPREADSHEET_ID")

@app.route("/")
def index():
    return render_template('upload.html')

@app.route("/detections", methods=["POST"])
def detections():
    try:
        image = request.files["file"]
        predictedFile = prediction.full_predict(image)
        return jsonify(predictedFile)
    except Exception as e:
        return jsonify({"detection error": str(e)}), 400

@app.route("/export", methods=["POST"])
def export():
    try:
        payload = request.get_json(force=True) or {}
        rows = payload.get("results", [])
        file_format = (payload.get("format") or "xlsx").lower()
 
        if file_format == "csv":
            headers = ["Meter Number", "Room Number"]
            output = io.StringIO()
            writer = csv.writer(output)
            writer.writerow(headers)
            for row in rows:
                writer.writerow([row.get("digit", ""), row.get("room", "")])
            
            mem = io.BytesIO(output.getvalue().encode("utf-8-sig"))
            mem.seek(0)
            return send_file(
                mem,
                mimetype="text/csv",
                as_attachment=True,
                download_name="detections.csv"
            )
        else:
            headers = ["Image", "File Name", "Meter Number", "Room Number"]
            wb = Workbook()
            ws = wb.active
            ws.title = "Detections"
            ws.append(headers)

            ws.column_dimensions['A'].width = 18 
            
            for idx, row in enumerate(rows, start=2): 
                ws.cell(row=idx, column=2, value=row.get("file_name", ""))
                ws.cell(row=idx, column=3, value=row.get("digit", ""))
                ws.cell(row=idx, column=4, value=row.get("room", ""))

                raw_img = row.get("predicted_img")
                if raw_img:
                    try:
                        if "base64," in raw_img:
                            clean_img = raw_img.split("base64,")[1]
                        else:
                            clean_img = raw_img
                        
                        img_bytes = base64.b64decode(clean_img)
                        img_stream = io.BytesIO(img_bytes)
                        
                        img = ExcelImage(img_stream)
                        img.width = 100
                        img.height = 100
                        
                        ws.row_dimensions[idx].height = 80 
                        ws.add_image(img, f'A{idx}')
                    except Exception as e:
                        print(f"แปลงรูปไม่สำเร็จในแถวที่ {idx}:", e)
            mem = io.BytesIO()
            wb.save(mem)
            mem.seek(0)
            return send_file(
                mem,
                mimetype="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
                as_attachment=True,
                download_name="detections.xlsx"
            )
    except Exception as e:
        return jsonify({"export error": str(e)}), 400

@app.route("/export_sheets", methods=["POST"])
def export_sheets():
    try:
        payload = request.get_json(force=True) or {}
        rows = payload.get("results", [])
        
        
        gc = gspread.service_account(filename='credentials.json')
        
        sh = gc.open_by_key(SHEET_ID)
        worksheet = sh.sheet1
        
        data_to_append = []
        current_time = datetime.now().strftime("%Y-%m-%d %H:%M:%S")
        
        for row in rows:
            data_to_append.append([
                current_time,
                row.get("file_name", ""),
                row.get("digit", ""),
                row.get("room", ""),
                row.get("confidence", ""),
                row.get("status", ""),
                row.get("error_message", "")
            ])
            
        worksheet.append_rows(data_to_append)
        return jsonify({"message": "success"}), 200

    except Exception as e:
        print("Sheet Error:", e)
        return jsonify({"error": str(e)}), 400


if __name__ == '__main__':
    app.run(debug=True)